"""``Bytehound --selftest-finish``: segment finishing in this build.

Records a short session with the bundled template config, closes it, and
checks that finishing ran in the helper process and produced complete,
readable output. Meant for verifying a frozen (PyInstaller) build, where the
helper is the application executable started in multiprocessing mode.
"""

from __future__ import annotations

import logging
import os
import sys
import tempfile
import time
from pathlib import Path

_LOG = logging.getLogger("bytehound.selftest")


def selftest_finish(rows: int = 600) -> int:
    from ..decoder.config_loader import load_config
    from ..decoder.frame_decoder import decode_frame
    from . import _isolated
    from .session import SessionLog, SessionSettings, read_manifest

    def say(msg: str) -> None:
        # A windowed (frozen) build has no stdout; the log file always works.
        _LOG.info("selftest-finish: %s", msg)
        if sys.stdout is not None:
            print(msg)

    failures = []
    template = Path(__file__).resolve().parents[1] / "resources" / "config_template"
    config = load_config(template)
    helper_pid = _isolated.run_isolated(os.getpid)
    if helper_pid == os.getpid():
        failures.append("finishing ran in the application process, not a helper")
    say(f"helper process pid {helper_pid} (application {os.getpid()})")

    folder = Path(tempfile.mkdtemp(prefix="bh_selftest_")) / "session"
    session = SessionLog(folder, config, SessionSettings(segment_seconds=3600, fixed_step_s=0.5),
                         config_source=template)
    session.start()
    frames = [decode_frame(config, fid, bytes((i * 7 + fid) & 0xFF for i in range(f.payload_length)), {})
              for fid, f in config.frames.items()]
    t = time.perf_counter()
    for i in range(rows):
        for f in frames:
            session.log_frame(f, t + i * 0.01)
    session.stop()
    t0 = time.monotonic()
    if not session.await_finished(300):
        failures.append("segment finishing did not complete within 300 s")
    say(f"segment finished in {time.monotonic() - t0:.1f}s")

    seg = read_manifest(folder)["segments"][0]
    if seg["status"] != "closed":
        failures.append(f"segment status {seg['status']}: {seg.get('errors')}")
    for key in ("event_xlsx", "fixed_xlsx", "event_analysis"):
        name = seg["files"].get(key)
        if not name or not (folder / name).is_file() or (folder / name).stat().st_size == 0:
            failures.append(f"{key} missing")
    if list(folder.glob("*.partial")):
        failures.append("temporary .partial files left behind")
    try:
        from openpyxl import load_workbook
        wb = load_workbook(folder / seg["files"]["event_xlsx"], read_only=True)
        data_rows = sum(1 for _ in wb["Data"].iter_rows()) - 1   # minus the header
        wb.close()
        if data_rows != seg["rows"]["event"]:
            failures.append(f"xlsx has {data_rows} rows, manifest says {seg['rows']['event']}")
        say(f"event xlsx readable: {data_rows} rows")
    except Exception as exc:
        failures.append(f"event xlsx unreadable: {exc}")

    _isolated.shutdown()
    if failures:
        for f in failures:
            say(f"FAIL: {f}")
        return 1
    say(f"OK: session finished in a helper process ({folder})")
    return 0
