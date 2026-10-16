"""save_workbook_atomic: an interrupted save never leaves a broken .xlsx."""

from __future__ import annotations

import pytest
from openpyxl import Workbook, load_workbook

from app.serial_logging._atomic import save_workbook_atomic


def _wb(value: str) -> Workbook:
    wb = Workbook()
    wb.active["A1"] = value
    return wb


def test_saves_and_replaces(tmp_path):
    target = tmp_path / "log.xlsx"
    save_workbook_atomic(_wb("first"), target)
    save_workbook_atomic(_wb("second"), target)
    assert load_workbook(target).active["A1"].value == "second"
    assert not (tmp_path / "log.xlsx.partial").exists()


def test_failed_save_keeps_previous_file(tmp_path):
    target = tmp_path / "log.xlsx"
    save_workbook_atomic(_wb("good"), target)

    class Exploding:
        def save(self, path):
            with open(path, "wb") as fp:
                fp.write(b"PK\x03\x04 truncated")
            raise OSError("disk vanished mid-save")

    with pytest.raises(OSError):
        save_workbook_atomic(Exploding(), target)
    assert load_workbook(target).active["A1"].value == "good"
    assert not (tmp_path / "log.xlsx.partial").exists()
