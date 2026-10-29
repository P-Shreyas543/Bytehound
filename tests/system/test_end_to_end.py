"""System tests: the complete app against a simulated BMS over a real wire.

Each test drives the real MainWindow (real config loader, PollingWorker
thread, parser, decoder, 60 Hz UI pipeline, loggers) the way an operator
would, and checks what the operator sees or what lands on disk / the wire.
"""

from __future__ import annotations

import os
import struct
import subprocess
import sys

import openpyxl
import pytest

from tests.system.conftest import CANONICAL_CONFIG, number, table_values, tcp_settings, wait_until
from tests.system.sim_device import FRAME_CELLS, PtyBms, encode_frame

ROOT = CANONICAL_CONFIG.parents[2]


def connect(window, settings) -> None:
    assert window._attempt_connect(settings), "connection failed"
    wait_until(lambda: window._packet_count >= 10, 5, "first telemetry")


def test_live_telemetry_is_decoded_into_the_table(window, device, error_boxes):
    connect(window, tcp_settings(device))

    values = table_values(window)
    assert number(values["Cell Voltage 1"]) == pytest.approx(3.300)
    assert number(values["Cell Voltage 2"]) == pytest.approx(3.310)
    assert values["BMS State"] == "Charge"                 # enum label
    assert values["FET Status"] == "Main FET Status"       # bitfield label
    # calc_groups "Cells": min|max|diff|avg
    assert number(values["Cells min"]) == pytest.approx(3.300)
    assert number(values["Cells max"]) == pytest.approx(3.310)
    assert number(values["Cells diff"]) == pytest.approx(0.010)
    assert number(values["Cells avg"]) == pytest.approx(3.305)
    assert error_boxes == []


def test_device_value_changes_reach_the_table(window, device):
    connect(window, tcp_settings(device))
    device.cell_mv = [3650, 3100]
    wait_until(lambda: number(table_values(window)["Cell Voltage 1"]) == pytest.approx(3.650),
               3, "updated cell voltage")
    assert number(table_values(window)["Cells diff"]) == pytest.approx(0.550)


def test_stream_recovers_from_garbage_and_corrupted_frames(window, device):
    connect(window, tcp_settings(device))
    bad_crc = bytearray(encode_frame(FRAME_CELLS, struct.pack(">HH", 1, 1)))
    bad_crc[-1] ^= 0xFF
    device.inject(b"\x00\xFF\xAA\x13\x55" + bytes(bad_crc) + b"\xAA\x55\x00")
    device.cell_mv = [3412, 3413]

    wait_until(lambda: number(table_values(window)["Cell Voltage 1"]) == pytest.approx(3.412),
               3, "a good frame after the corrupted bytes")
    # The corrupted frame's values (0.001 V) must never have been displayed.
    assert number(table_values(window)["Cells min"]) == pytest.approx(3.412)


def test_tx_command_from_the_panel_arrives_on_the_wire(window, device):
    connect(window, tcp_settings(device))

    window._tx_command_combo.setCurrentText("Set Charge Current")
    window._tx_field_inputs["Current Limit"].setText("12.5")   # A, scale 0.1
    window._send_tx_command()

    expected = encode_frame(0x8002, struct.pack(">H", 125))
    wait_until(lambda: expected in device.received(), 3, "Set Charge Current frame")

    window._tx_command_combo.setCurrentText("Request Fault Codes")
    window._send_tx_command()
    wait_until(lambda: encode_frame(0x8001, b"\x01") in device.received(), 3, "Request Fault Codes frame")


@pytest.mark.parametrize("console", ["console shown", "console hidden"])
def test_logging_session_writes_raw_and_decoded_workbooks(window, device, tmp_path, console):
    # The raw log used to be fed only while the console dock was visible:
    # with it hidden (or tabbed away) the raw log came out empty.
    window._console_dock.setVisible(console == "console shown")
    connect(window, tcp_settings(device))
    assert window._start_logging_file(tmp_path / "session.xlsx", choice="Raw + Decoded")
    window._tx_command_combo.setCurrentText("Request Fault Codes")
    window._send_tx_command()
    # A long window, so the few packets already queued in the UI when
    # logging started (counted but legitimately not logged) don't matter.
    start = window._packet_count
    wait_until(lambda: window._packet_count >= start + 200, 10, "logged telemetry")
    received = window._packet_count - start
    window._stop_logging()

    raw_path, decoded_path = tmp_path / "session_raw.xlsx", tmp_path / "session_decoded.xlsx"
    assert raw_path.exists() and decoded_path.exists()

    raw = openpyxl.load_workbook(raw_path, read_only=True)["Data"]
    rows = list(raw.iter_rows(values_only=True))[1:]   # timestamp, direction, hex, delta_t_ms
    rx = [r for r in rows if r[1] == "RX"]
    tx = [r for r in rows if r[1] == "TX"]
    assert len(rx) >= received * 0.9, f"raw log kept {len(rx)} of {received} received packets"
    assert all(r[2].replace(" ", "").upper().startswith("AA55") for r in rx)
    assert any(r[2].replace(" ", "").upper() == encode_frame(0x8001, b"\x01").hex().upper() for r in tx), \
        "TX command missing from raw log"

    wb = openpyxl.load_workbook(decoded_path, read_only=True)
    assert wb.sheetnames == ["Metadata", "Data"]
    rows = list(wb["Data"].iter_rows(values_only=True))
    header = [str(h) for h in rows[0]]
    volts = header.index("0x10.Cell Voltage 1 (V)")
    state = header.index("0x30.BMS State.label")
    # One wide row per telemetry cycle (both frames).
    assert len(rows) - 1 >= received // 2 * 0.9
    assert all(r[volts] == pytest.approx(3.300) for r in rows[1:])
    assert all(r[state] == "Charge" for r in rows[1:])


def test_auto_reconnect_after_the_device_drops(window, device):
    connect(window, tcp_settings(device, auto_reconnect=True))
    assert device.connections == 1

    device.drop_connection()
    wait_until(lambda: device.connections == 2, 8, "automatic reconnection")
    device.cell_mv = [3555, 3556]
    wait_until(lambda: number(table_values(window)["Cell Voltage 1"]) == pytest.approx(3.555),
               5, "telemetry after reconnect")


def test_disconnect_stops_the_worker_and_releases_the_port(window, device):
    connect(window, tcp_settings(device))
    worker = window._serial
    window._disconnect()
    assert window._serial is None
    assert not worker.isRunning()
    # The device can accept a fresh session afterwards.
    connect(window, tcp_settings(device))
    assert device.connections == 2


@pytest.mark.skipif(os.name != "posix", reason="pseudo-terminals are POSIX only")
def test_serial_port_end_to_end_over_a_pty(window):
    from app.serial_io.serial_worker import SerialSettings

    bms = PtyBms().start()
    try:
        connect(window, SerialSettings(port=bms.port, baud_rate=115200))
        assert number(table_values(window)["Cell Voltage 2"]) == pytest.approx(3.310)
        window._tx_command_combo.setCurrentText("Request Fault Codes")
        window._send_tx_command()
        wait_until(lambda: encode_frame(0x8001, b"\x01") in bms.received(), 3, "TX over serial")
    finally:
        window._disconnect()
        bms.stop()


@pytest.mark.parametrize("config, expected_rc", [
    (CANONICAL_CONFIG, 0),
    (ROOT / "app" / "resources" / "sample_raw_log.txt", 1),   # not a config
])
def test_config_validator_cli(config, expected_rc):
    proc = subprocess.run(
        [sys.executable, "-m", "app.main", "--validate", str(config)],
        cwd=ROOT, capture_output=True, text=True, timeout=120,
        env={**os.environ, "QT_QPA_PLATFORM": "offscreen"},
    )
    assert proc.returncode == expected_rc, proc.stdout + proc.stderr
