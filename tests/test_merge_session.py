"""Session merge (app/serial_logging/session_merge.py, Tools → Merge Session…,
tools/merge_session.py): a session's hourly segments joined into one
workbook, or one per day, without touching the session folder."""

from __future__ import annotations

import csv
import time
from datetime import datetime, timedelta
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.decoder.config_loader import load_config
from app.decoder.frame_decoder import decode_frame
from app.serial_logging.session import SessionLog, SessionSettings, read_manifest
from app.serial_logging.session_merge import MergeCancelled, MergeError, merge, merge_job, session_summary
from tools.merge_session import main as merge_cli

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "app" / "resources" / "config_template"
FIXTURE = ROOT / "tests" / "fixtures" / "canonical_config"


class Clock:
    def __init__(self, wall: float) -> None:
        self.wall = wall
        self.perf = 5_000.0

    def advance(self, s: float) -> None:
        self.wall += s
        self.perf += s


def _frames(config):
    out = []
    for fid, f in config.frames.items():
        n = f.payload_length or max((s.end_byte for s in config.signals_by_frame.get(fid, [])), default=0)
        out.append(decode_frame(config, fid, bytes((i * 7 + fid) & 0xFF for i in range(n)), {}))
    return out


def _record(folder: Path, start_wall: float, seconds: float, *, fixed: bool = False, change_at: float = None):
    config = load_config(TEMPLATE)
    clock = Clock(start_wall)
    s = SessionLog(folder, config, SessionSettings(segment_seconds=600, fixed_step_s=1.0 if fixed else None),
                   config_source=TEMPLATE, clock=lambda: clock.wall, perf=lambda: clock.perf,
                   disk_free=lambda p: 500 * 1024 ** 3)
    s.start()
    t = 0.0
    frames = _frames(config)
    while t < seconds:
        if change_at is not None and t >= change_at:
            s.connection_lost("configuration change")
            config = load_config(FIXTURE)
            s.change_config(config, FIXTURE)
            s.connection_restored()
            frames = _frames(config)
            change_at = None
        for f in frames:
            s.log_frame(f, clock.perf)
        clock.advance(5.0)
        t += 5.0
        s.tick()
    s.stop()
    assert s.await_finished(300)
    return read_manifest(folder)


def _xlsx_rows(path: Path):
    wb = load_workbook(path, read_only=True)
    data = [r for name in wb.sheetnames if name.startswith("Data") for r in list(wb[name].iter_rows(values_only=True))[1:]]
    meta = {r[0]: r[1] for r in list(wb["Metadata"].iter_rows(values_only=True))[1:]}
    header = next(wb["Data"].iter_rows(values_only=True))
    wb.close()
    return header, data, meta


def _local_midnight_ahead() -> float:
    now = datetime.now()
    midnight = datetime(now.year, now.month, now.day) + timedelta(days=1)
    return time.mktime(midnight.timetuple())


def test_one_workbook_holds_every_row_in_order(tmp_path):
    folder = tmp_path / "sess"
    m = _record(folder, time.time() - 7200, 1790)
    before = (folder / "session.json").read_bytes()
    out = merge(folder)
    assert len(out) == 1
    total = sum(s["rows"]["event"] for s in m["segments"])
    header, data, meta = _xlsx_rows(Path(out[0]["xlsx"]))
    assert len(m["segments"]) == 3 and len(data) == total == out[0]["rows"]
    with (folder / m["segments"][0]["files"]["event_csv"]).open(encoding="utf-8-sig") as f:
        assert list(header) == next(csv.reader(f))                 # same columns as the segments
    ti = header.index("time_utc")
    times = [r[ti] for r in data]
    assert times == sorted(times)                                    # segments joined in order
    assert meta["session"] == "sess" and meta["rows"] == str(total) and meta["segments"] == "1..3 (3)"
    assert (folder / "session.json").read_bytes() == before         # session folder untouched


def test_split_by_day_cuts_exactly_at_local_midnight(tmp_path):
    folder = tmp_path / "sess"
    m = _record(folder, _local_midnight_ahead() - 1500, 3000)       # 25 min before .. 25 min after
    out = merge(folder, split="day")
    assert len(out) == 2
    total = sum(s["rows"]["event"] for s in m["segments"])
    assert sum(w["rows"] for w in out) == total
    for w in out:
        day = Path(w["xlsx"]).stem.rsplit("_", 1)[-1]
        header, data, _meta = _xlsx_rows(Path(w["xlsx"]))
        ti = header.index("time_local")
        assert data and all(str(r[ti]).startswith(day) for r in data)


def test_a_configuration_change_starts_a_new_part(tmp_path):
    folder = tmp_path / "sess"
    m = _record(folder, time.time() - 7200, 1500, change_at=700)
    assert [s["config"] for s in m["segments"]][-1] == "config_2"
    out = merge(folder, keep_csv=True)
    assert len(out) == 2 and Path(out[1]["xlsx"]).stem.endswith("_part2")
    h1, d1, _ = _xlsx_rows(Path(out[0]["xlsx"]))
    h2, d2, _ = _xlsx_rows(Path(out[1]["xlsx"]))
    assert h1 != h2 and len(d1) + len(d2) == sum(s["rows"]["event"] for s in m["segments"])
    assert all(Path(w["csv"]).exists() for w in out)                      # --csv keeps the merged CSV


def test_fixed_step_log_merges_and_missing_kind_is_explained(tmp_path):
    folder = tmp_path / "sess"
    m = _record(folder, time.time() - 7200, 1300, fixed=True)
    out = merge(folder, kind="fixed")
    _h, data, _ = _xlsx_rows(Path(out[0]["xlsx"]))
    assert len(data) == sum(s["rows"]["fixed"] for s in m["segments"])
    plain = tmp_path / "plain"
    _record(plain, time.time() - 7200, 100)
    with pytest.raises(MergeError, match="no fixed log"):
        merge(plain, kind="fixed")


def test_a_recording_session_is_refused(tmp_path):
    folder = tmp_path / "sess"
    _record(folder, time.time() - 7200, 100)
    import json
    m = json.loads((folder / "session.json").read_text(encoding="utf-8"))
    m["status"] = "recording"
    (folder / "session.json").write_text(json.dumps(m), encoding="utf-8")
    with pytest.raises(MergeError, match="still recording"):
        merge(folder)


def test_cancel_leaves_nothing_behind(tmp_path):
    folder = tmp_path / "sess"
    _record(folder, time.time() - 7200, 1790)
    out = tmp_path / "out"
    calls = {"n": 0}

    def cancel():
        calls["n"] += 1
        return calls["n"] > 1                  # cancel after the first segment

    with pytest.raises(MergeCancelled):
        merge(folder, out=out, cancel=cancel)
    assert list(out.iterdir()) == []


def test_progress_reports_rows(tmp_path):
    folder = tmp_path / "sess"
    m = _record(folder, time.time() - 7200, 1790)
    seen = []
    merge(folder, progress=lambda phase, done, total: seen.append((phase, done, total)))
    total = sum(s["rows"]["event"] for s in m["segments"])
    assert seen and all(t in (total, seen[-1][2]) for _p, _d, t in seen)
    assert any(p.startswith("Writing workbook 1 of 1") for p, _d, _t in seen)


def test_summary_and_child_process_job(tmp_path):
    """What the Tools dialog shows, and the job it runs in a child process."""
    import multiprocessing as mp

    folder = tmp_path / "sess"
    m = _record(folder, time.time() - 7200, 1790, fixed=True)
    info = session_summary(folder)
    assert info["segments"] == len(m["segments"]) and set(info["rows"]) == {"event", "fixed"}
    with pytest.raises(MergeError, match="not a session folder"):
        session_summary(tmp_path)
    ctx = mp.get_context("spawn")
    q, ev = ctx.Queue(), ctx.Event()
    proc = ctx.Process(target=merge_job, args=(str(folder), "event", "none", str(tmp_path / "o"), False, q, ev))
    proc.start()
    msgs = []
    while True:
        msg = q.get(timeout=120)
        msgs.append(msg)
        if msg[0] in ("done", "error", "cancelled"):
            break
    proc.join(30)
    assert msgs[-1][0] == "done", msgs[-1]
    assert Path(msgs[-1][1][0]["xlsx"]).exists()


def test_command_line(tmp_path, capsys):
    folder = tmp_path / "sess"
    _record(folder, time.time() - 7200, 100)
    assert merge_cli([str(folder), "--out", str(tmp_path / "o")]) == 0
    assert "rows=" in capsys.readouterr().out
    assert merge_cli([str(tmp_path)]) == 1                         # not a session: message, exit 1


def test_compiler_cancel_mid_write_writes_nothing(tmp_path):
    from app.serial_logging.decoded_logger import CompileCancelled, compile_csv_to_xlsx

    src = tmp_path / "big.csv"
    rows = "".join(f"{i},{i * 2}\n" for i in range(25_000))
    src.write_text("a,b\n" + rows, encoding="utf-8")
    seen = []
    exc = compile_csv_to_xlsx(src, {"k": "v"}, tmp_path / "out.xlsx", progress=seen.append,
                              cancel=lambda: len(seen) >= 1)
    assert isinstance(exc, CompileCancelled) and seen == [10_000]
    assert not (tmp_path / "out.xlsx").exists() and not list(tmp_path.glob("*.partial"))


def test_selftest_merge_passes_in_this_environment():
    """The check the frozen build runs (Bytehound.exe --selftest-merge)."""
    from app.serial_logging.selftest import selftest_merge

    assert selftest_merge(rows=200) == 0
