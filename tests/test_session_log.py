"""Long-duration logging sessions: segments, manifest, fixed-step grid,
gaps, disk guard and crash recovery (Plane.md Part A, Part H, F9–F11)."""

from __future__ import annotations

import csv
import json
import time
from pathlib import Path
from typing import Any, Dict, List, Optional

import pytest
from openpyxl import load_workbook

from app.decoder.config_loader import load_config
from app.decoder.frame_decoder import DecodedFrame, DecodedSignal, decode_frame
from app.serial_logging import decoded_logger as dl_mod
from app.serial_logging import session as session_mod
from app.serial_logging.session import (
    FixedStepSampler,
    SessionLog,
    SessionSettings,
    find_unfinished_sessions,
    read_manifest,
    recover_session,
)

TEMPLATE = Path(__file__).resolve().parents[1] / "app" / "resources" / "config_template"


@pytest.fixture(scope="module")
def config():
    return load_config(TEMPLATE)


class Clock:
    """Controllable wall + monotonic clocks."""

    def __init__(self) -> None:
        self.wall = 1_800_000_000.0
        self.perf = 5_000.0

    def advance(self, s: float) -> None:
        self.wall += s
        self.perf += s


class RawHooks:
    def __init__(self) -> None:
        self.started: List[Path] = []
        self.stopped = 0

    def start(self, path: Path, metadata: Dict[str, Any]) -> None:
        self.started.append(path)
        path.write_bytes(b"")  # the real worker creates the .bhraw file

    def stop(self) -> Optional[str]:
        self.stopped += 1
        return None


def _frames(config) -> List[DecodedFrame]:
    out = []
    for fid, f in config.frames.items():
        payload = bytes((i * 7 + fid) & 0xFF for i in range(f.payload_length))
        out.append(decode_frame(config, fid, payload, {}))
    return out


def _session(tmp_path, config, clock, **settings_kw) -> tuple[SessionLog, RawHooks]:
    hooks = RawHooks()
    s = SessionLog(
        tmp_path / "sess", config, SessionSettings(**settings_kw),
        config_source=TEMPLATE, metadata={"app": "test"}, app_version="9.9.9",
        raw_capture=hooks, clock=lambda: clock.wall, perf=lambda: clock.perf,
        disk_free=lambda p: 500 * 1024 ** 3,
    )
    return s, hooks


def _csv_rows(path: Path) -> List[List[str]]:
    with path.open(newline="", encoding="utf-8") as fp:
        return list(csv.reader(fp))


# ---------------------------------------------------------------------------
# Fixed-step sampler
# ---------------------------------------------------------------------------

def _sig(fid: int, name: str, value: float) -> DecodedFrame:
    return DecodedFrame(fid, "F", [DecodedSignal(fid, "F", name, value, value, "V", "ok")])


def _sampler(step=0.1, hold=True):
    cols = [((1, "a"), "a (V)"), ((2, "b"), "b (V)")]
    return FixedStepSampler(cols, step, 100.0, stale_steps=3, hold_stale=hold, wall_of=lambda t: 1e9 + t)


def test_sampler_rows_hold_newest_value_at_or_before_each_tick():
    s = _sampler()
    rows = s.feed(_sig(1, "a", 1.0), 100.05)       # tick 0 complete; a not seen yet
    assert [r[0] for r in rows] == [0] and rows[0][4:8] == ["", "", "", ""] and rows[0][-1] == 2
    rows = s.feed(_sig(1, "a", 2.0), 100.20)                         # completes 0.1 (and not 0.2)
    assert [r[0] for r in rows] == [1]
    assert rows[0][4:6] == [1.0, 50]                                  # a=1.0, 50 ms old at t=0.1
    rows = s.feed(_sig(2, "b", 7.0), 100.31)
    assert [r[0] for r in rows] == [2, 3]
    assert rows[0][4:6] == [2.0, 0]                                   # frame exactly at the tick counts
    assert rows[1][6:8] == ["", ""] and rows[1][-1] == 1             # b not yet seen → stale


def test_sampler_marks_and_optionally_blanks_stale_values():
    for hold, expect in ((True, 1.0), (False, "")):
        s = _sampler(hold=hold)
        s.feed(_sig(1, "a", 1.0), 100.0)
        rows = s.advance(100.5)                    # a is 500 ms old at t=0.5 > 3 steps
        last = rows[-1]
        assert last[0] == 5 and last[5] == 500
        assert last[4] == expect
        assert last[-1] == 2                       # a stale + b never seen


def test_sampler_catches_up_after_a_stall_without_skipping():
    s = _sampler()
    s.feed(_sig(1, "a", 1.0), 100.0)
    rows = s.feed(_sig(1, "a", 2.0), 160.0)        # GUI stalled for a minute
    assert [r[0] for r in rows] == list(range(0, 600))
    assert all(r[4] == 1.0 for r in rows)          # every row holds the value valid at its tick


def test_sampler_grid_does_not_drift_over_a_week():
    s = _sampler(step=0.1)
    n = 7 * 24 * 3600 * 10                          # one week of 0.1 s ticks
    assert abs(s.grid_time(n) - (100.0 + 604_800.0)) < 1e-6


# ---------------------------------------------------------------------------
# Session folder, segments and manifest
# ---------------------------------------------------------------------------

def test_session_rotates_segments_and_records_everything(tmp_path, config):
    clock = Clock()
    s, hooks = _session(tmp_path, config, clock, segment_seconds=60, fixed_step_s=1.0)
    s.start()
    frames = _frames(config)
    for _seg in range(3):
        for _ in range(30):                          # 30 s per loop, 2 loops per segment
            for f in frames:
                s.log_frame(f, clock.perf)
            clock.advance(1.0)
            s.tick()
        s.log_rejected(b"\xAA\x55\x01", 0x1000, "CRC mismatch", clock.perf)
        clock.advance(30.0)
        s.tick()
    s.stop()
    assert s.await_finished(30)

    m = read_manifest(tmp_path / "sess")
    assert m["status"] == "closed" and m["app_version"] == "9.9.9"
    assert m["config"]["files_sha256"] and (tmp_path / "sess" / "config" / "frames.csv").exists()
    segs = m["segments"]
    assert [x["index"] for x in segs] == list(range(1, len(segs) + 1)) and len(segs) >= 3
    assert all(x["status"] == "closed" for x in segs), segs
    assert [p.name for p in hooks.started] == [f"seg_{i:04d}.bhraw" for i in range(1, len(segs) + 1)]
    assert hooks.stopped == len(segs)
    folder = tmp_path / "sess"
    for x in segs:
        for name in ("event_csv", "event_xlsx", "fixed_csv", "fixed_xlsx"):
            assert (folder / x["files"][name]).exists()
        assert all(len(d) == 64 for d in x["sha256"].values())
    total_event = sum(x["rows"]["event"] for x in segs)
    assert total_event > 0
    assert sum(x.get("rejected_frames", 0) for x in segs) == 3
    first = _csv_rows(folder / segs[0]["files"]["event_csv"])
    assert first[0][:2] == ["time_utc", "time_local"]
    assert first[1][0].endswith("+00:00") and "T" in first[1][0]
    wb = load_workbook(folder / segs[0]["files"]["event_xlsx"], read_only=True)
    assert wb.sheetnames == ["Metadata", "Data"]
    wb.close()


def test_fixed_step_grid_is_continuous_across_segments(tmp_path, config):
    clock = Clock()
    s, _ = _session(tmp_path, config, clock, segment_seconds=10, fixed_step_s=0.5, event_log=False)
    s.start()
    frames = _frames(config)
    for _ in range(100):                             # 25 s of data at 4 Hz
        s.log_frame(frames[0], clock.perf)
        clock.advance(0.25)
        s.tick()
    clock.advance(2.0)
    s.tick()
    s.stop()
    assert s.await_finished(30)
    m = read_manifest(tmp_path / "sess")
    ticks: List[int] = []
    for x in m["segments"]:
        rows = _csv_rows(tmp_path / "sess" / x["files"]["fixed_csv"])
        ticks += [int(r[0]) for r in rows[1:]]
    assert len(m["segments"]) >= 3
    assert ticks == list(range(len(ticks)))          # no gaps, no repeats across rotations
    assert len(ticks) >= 50


def test_segments_rotate_before_the_xlsx_row_limit(tmp_path, config):
    clock = Clock()
    s, _ = _session(tmp_path, config, clock, segment_seconds=10_000, max_segment_rows=40)
    s.start()
    frames = _frames(config)
    for _ in range(100):
        for f in frames:
            s.log_frame(f, clock.perf)
        clock.advance(0.1)
        time.sleep(0.002)  # let the writer thread count the rows it wrote
        s.tick()
    s.stop()
    assert s.await_finished(30)
    segs = read_manifest(tmp_path / "sess")["segments"]
    assert len(segs) >= 2
    assert all(x["rows"]["event"] <= 40 + 10 for x in segs)  # rotation granularity: one tick


def test_connection_gap_closes_segment_and_session_continues(tmp_path, config):
    clock = Clock()
    s, hooks = _session(tmp_path, config, clock, fixed_step_s=1.0)
    s.start()
    f = _frames(config)[0]
    s.log_frame(f, clock.perf)
    clock.advance(5)
    s.tick()
    s.connection_lost("USB device disconnected")
    clock.advance(42)
    s.tick()                                          # no segment open: nothing written, no error
    s.connection_restored()
    s.log_frame(f, clock.perf)
    clock.advance(3)
    s.tick()
    s.stop()
    assert s.await_finished(30)
    m = read_manifest(tmp_path / "sess")
    assert len(m["segments"]) == 2
    assert m["gaps"][0]["reason"] == "USB device disconnected"
    assert m["gaps"][0]["seconds"] == 42.0
    rows = _csv_rows(tmp_path / "sess" / m["segments"][1]["files"]["fixed_csv"])
    stale_counts = [int(r[-1]) for r in rows[1:]]
    assert len(rows) - 1 >= 40                        # gap ticks are written, not skipped
    assert max(stale_counts) > 0                      # and marked stale


def test_disk_guard_warns_then_stops(tmp_path, config):
    clock = Clock()
    warnings: List[str] = []
    errors: List[str] = []
    free = {"bytes": 50 * 1024 ** 3}
    s = SessionLog(
        tmp_path / "sess", config, SessionSettings(), raw_capture=RawHooks(),
        on_warning=warnings.append, on_error=errors.append,
        clock=lambda: clock.wall, perf=lambda: clock.perf, disk_free=lambda p: free["bytes"],
    )
    s.start()
    s.tick()
    assert not warnings and not errors
    free["bytes"] = 5 * 1024 ** 3
    clock.advance(31)
    s.tick()
    assert len(warnings) == 1 and "getting full" in warnings[0]
    clock.advance(31)
    s.tick()
    assert len(warnings) == 1                          # warned once, not every check
    free["bytes"] = 1 * 1024 ** 3
    clock.advance(31)
    s.tick()
    assert s.stopped_for_disk and "Logging stopped" in errors[0]
    s.stop()
    assert s.await_finished(30)
    texts = [e["text"] for e in read_manifest(tmp_path / "sess")["events"]]
    assert any("low" in t for t in texts) and any("stopped" in t for t in texts)


def test_existing_session_is_never_overwritten(tmp_path, config):
    clock = Clock()
    s, _ = _session(tmp_path, config, clock)
    s.start()
    s.stop()
    assert s.await_finished(30)
    again, _ = _session(tmp_path, config, clock)
    with pytest.raises(FileExistsError):
        again.start()


def test_recovery_after_a_crash_keeps_all_rows(tmp_path, config, monkeypatch):
    clock = Clock()
    s, _ = _session(tmp_path, config, clock, fixed_step_s=1.0)
    s.start()
    frames = _frames(config)
    for _ in range(20):
        for f in frames:
            s.log_frame(f, clock.perf)
        clock.advance(1.0)
        s.tick()
    # Simulate the process dying: writers flush their temp CSVs but nothing
    # is compiled and the manifest still says "recording".
    monkeypatch.setattr(dl_mod, "finalize_temp_log", lambda *a, **k: None)
    monkeypatch.setattr(session_mod, "finalize_temp_log", lambda *a, **k: None)
    seg = s._segment
    seg.event.close(timeout=30)
    seg.fixed.close()
    seg.fixed.join(30)
    event_rows = seg.event.rows_written
    fixed_rows = seg.fixed.rows_written
    monkeypatch.undo()

    folder = tmp_path / "sess"
    assert read_manifest(folder)["status"] == "recording"
    assert find_unfinished_sessions(tmp_path) == [folder]
    assert recover_session(folder)
    m = read_manifest(folder)
    assert m["status"] == "recovered" and m["segments"][0]["status"] == "recovered"
    assert len(_csv_rows(folder / "seg_0001_event.csv")) - 1 == event_rows
    assert len(_csv_rows(folder / "seg_0001_fixed.csv")) - 1 == fixed_rows
    assert (folder / "seg_0001_event.xlsx").exists() and (folder / "seg_0001_fixed.xlsx").exists()
    assert not list(folder.glob("*.tmp_data"))
    assert find_unfinished_sessions(tmp_path) == []
    assert recover_session(folder) is False            # idempotent
    json.loads((folder / "session.json").read_text(encoding="utf-8"))


def test_fixed_step_values_match_the_event_log_at_each_tick(tmp_path, config):
    """At every tick the fixed-step value equals the newest event-log value
    received at or before that instant (Plane.md Part H)."""

    clock = Clock()
    s, _ = _session(tmp_path, config, clock, fixed_step_s=1.0)
    s.start()
    fid = next(iter(config.frames))
    length = config.frames[fid].payload_length
    sent = []                                          # (elapsed_s, decoded) as logged
    for i in range(40):
        clock.advance(0.35 if i % 3 else 0.6)          # irregular arrivals, never on a tick
        payload = bytes([(i * 13) & 0xFF] * length)
        d = decode_frame(config, fid, payload, {})
        s.log_frame(d, clock.perf)
        sent.append((clock.perf - s.session_perf0, d))
        s.tick()
    clock.advance(2.0)
    s.tick()
    s.stop()
    assert s.await_finished(30)

    folder = tmp_path / "sess"
    fixed = _csv_rows(folder / "seg_0001_fixed.csv")
    header = fixed[0]
    sig = next(sig for sig in sent[0][1].signals if sig.status == "ok")
    col = next(i for i, h in enumerate(header) if h.startswith(f"0x{fid:X}.{sig.signal_name}"))
    checked = 0
    for row in fixed[1:]:
        t = float(row[3])
        before = [d for (at, d) in sent if at <= t]
        if not before:
            assert row[col] == ""
            continue
        expect = next(x for x in before[-1].signals if x.signal_name == sig.signal_name).scaled_value
        assert float(row[col]) == pytest.approx(expect)
        checked += 1
    assert checked >= 10


def test_whole_session_rebuild_from_raw_captures(tmp_path, config):
    """Tools → Rebuild with a session.json rebuilds every segment from its
    .bhraw, and the rebuilt event rows match what was logged live."""
    from app.protocol.packet_builder import build_packet
    from app.serial_logging.raw_capture import RX, RawCapture
    from app.serial_logging.session import rebuild_session_event_logs

    class RealRaw:
        cap = None

        def start(self, path, metadata):
            self.cap = RawCapture(path, metadata)

        def stop(self):
            cap, self.cap = self.cap, None
            if cap is not None:
                cap.close()
                return cap.error
            return None

    clock = Clock()
    raw = RealRaw()
    s = SessionLog(
        tmp_path / "sess", config, SessionSettings(segment_seconds=30), raw_capture=raw,
        clock=lambda: clock.wall, perf=lambda: clock.perf, disk_free=lambda p: 1 << 40,
    )
    s.start()
    for i in range(70):
        for fid, f in config.frames.items():
            payload = bytes(((i + fid) & 0xFF,) * f.payload_length)
            pkt = build_packet(config.protocol, fid, payload)
            raw.cap.write(RX, clock.perf, pkt)
            s.log_frame(decode_frame(config, fid, payload, {}), clock.perf)
        clock.advance(1.0)
        s.tick()
    s.stop()
    assert s.await_finished(30)

    folder = tmp_path / "sess"
    live = read_manifest(folder)["segments"]
    assert len(live) >= 2
    assert rebuild_session_event_logs(folder, config) == len(live)
    m = read_manifest(folder)
    for seg in m["segments"]:
        wb = load_workbook(folder / seg["files"]["event_rebuilt_xlsx"], read_only=True)
        rebuilt_rows = sum(1 for _ in wb["Data"].iter_rows()) - 1
        wb.close()
        assert rebuilt_rows == seg["rows"]["event"]
    assert any("rebuilt" in e["text"] for e in m["events"])


def test_finished_segment_threads_are_released(tmp_path, config):
    """Retained finisher threads leaked ~2 OS handles per segment on Windows
    (found by tools/soak_session.py); only live finishers may be kept."""
    clock = Clock()
    s, _ = _session(tmp_path, config, clock, segment_seconds=5)
    s.start()
    f = _frames(config)[0]
    for _ in range(200):                     # ~40 rotations
        s.log_frame(f, clock.perf)
        clock.advance(1.0)
        s.tick()
        time.sleep(0.002)
    assert len(s._finishers) <= 5
    s.stop()
    assert s.await_finished(60)


def test_closed_session_has_complete_analysis_data(tmp_path, config):
    """Every closed segment gets an analysis cache, and the session summary
    already covers all segments when logging stops — so the Analysis Suite
    opens it without any preparation (Part B)."""
    import numpy as np

    clock = Clock()
    s, _ = _session(tmp_path, config, clock, segment_seconds=20, fixed_step_s=1.0)
    s.start()
    frames = _frames(config)
    for _ in range(70):
        for f in frames:
            s.log_frame(f, clock.perf)
        clock.advance(1.0)
        s.tick()
        time.sleep(0.002)
    s.stop()
    assert s.await_finished(60)
    folder = tmp_path / "sess"
    m = read_manifest(folder)
    indices = [x["index"] for x in m["segments"]]
    assert len(indices) >= 3
    for kind in ("event", "fixed"):
        assert all((folder / x["files"][f"{kind}_analysis"]).exists() for x in m["segments"])
        with np.load(folder / f"{kind}_summary.npz", allow_pickle=True) as z:
            assert [int(i) for i in z["segments"]] == indices


def test_segment_integrity_includes_driver_errors(tmp_path, config):
    """Per-segment integrity counters in session.json include the serial
    driver's error flags (they arrive as a nested dict)."""
    clock = Clock()
    s, _ = _session(tmp_path, config, clock)
    s.start()
    s.set_integrity({"frames_ok": 10, "discarded_bytes": 0,
                     "driver_errors": {"hardware_overrun": 1, "framing_error": 0}})
    clock.advance(5)
    s.set_integrity({"frames_ok": 50, "discarded_bytes": 3,
                     "driver_errors": {"hardware_overrun": 4, "framing_error": 2}})
    s.stop()
    assert s.await_finished(30)
    integ = read_manifest(tmp_path / "sess")["segments"][0]["integrity"]
    assert integ["frames_ok"] == 40 and integ["discarded_bytes"] == 3
    assert integ["driver_hardware_overrun"] == 3 and integ["driver_framing_error"] == 2


class RotatingRawHooks(RawHooks):
    def __init__(self) -> None:
        super().__init__()
        self.rotated: List[Path] = []

    def rotate(self, path: Path, metadata: Dict[str, Any]) -> Optional[str]:
        self.rotated.append(path)
        path.write_bytes(b"")
        return "old writer error" if len(self.rotated) == 1 else None


def test_rotation_switches_raw_capture_without_a_gap(tmp_path, config):
    """Found on the NUCLEO HIL rig: stop() then start() at every rotation
    left ~70 ms of received bytes out of every .bhraw file. With rotate()
    the next file starts before the current segment closes."""
    clock = Clock()
    hooks = RotatingRawHooks()
    s = SessionLog(tmp_path / "sess", config, SessionSettings(segment_seconds=10), raw_capture=hooks,
                   clock=lambda: clock.wall, perf=lambda: clock.perf, disk_free=lambda p: 1 << 40)
    s.start()
    frames = _frames(config)
    for _ in range(25):
        for f in frames:
            s.log_frame(f, clock.perf)
        clock.advance(1.0)
        s.tick()
    s.stop()
    assert s.await_finished(30)
    segs = read_manifest(tmp_path / "sess")["segments"]
    assert len(segs) == 3
    assert [p.name for p in hooks.started] == ["seg_0001.bhraw"]
    assert [p.name for p in hooks.rotated] == ["seg_0002.bhraw", "seg_0003.bhraw"]
    assert hooks.stopped == 1                         # only at the end of the session
    assert [x["files"]["raw"] for x in segs] == [f"seg_{i:04d}.bhraw" for i in (1, 2, 3)]
    assert segs[0]["raw_capture_error"] == "old writer error"
    assert "raw_capture_error" not in segs[1]


def test_segment_integrity_records_gauges_as_values(tmp_path, config):
    clock = Clock()
    s, _ = _session(tmp_path, config, clock, segment_seconds=10)
    s.start()
    s.set_integrity({"parser_bytes_in": 100, "buffered_bytes": 7, "rx_service_gap_max_ms": 12.5})
    clock.advance(11.0)
    s.set_integrity({"parser_bytes_in": 250, "buffered_bytes": 3, "rx_service_gap_max_ms": 40.0})
    s.tick()
    s.stop()
    assert s.await_finished(30)
    first = read_manifest(tmp_path / "sess")["segments"][0]["integrity"]
    assert first["parser_bytes_in"] == 150
    assert first["buffered_bytes"] == 3 and first["rx_service_gap_max_ms"] == 40.0


def test_helper_timeout_marks_the_segment_and_keeps_the_csv(tmp_path, config, monkeypatch):
    """task.md Phase 2: a hung xlsx compile must not be presented as a
    finished segment. The CSV stays, no .xlsx or .partial is left, and the
    reason is in session.json."""
    from app.serial_logging import _isolated, session_analysis
    from tests import isolated_jobs as jobs

    # Jobs that really hang: a tiny timeout on real jobs races with them
    # finishing, and a job that wins the race is (correctly) kept.
    monkeypatch.setattr(dl_mod, "compile_csv_to_xlsx", jobs.hang_compile)
    monkeypatch.setattr(session_analysis, "build_segment_cache", jobs.hang)
    monkeypatch.setattr(_isolated, "timeout_for", lambda *p: 1.0)
    clock = Clock()
    s, _ = _session(tmp_path, config, clock, segment_seconds=3600, fixed_step_s=1.0)
    s.start()
    for f in _frames(config) * 20:
        s.log_frame(f, clock.perf)
        clock.advance(0.1)
    s.stop()
    assert s.await_finished(120)
    seg = read_manifest(tmp_path / "sess")["segments"][0]
    assert seg["status"] == "closed_with_errors"
    assert any("did not finish within" in e for e in seg["errors"]), seg["errors"]
    folder = tmp_path / "sess"
    assert (folder / seg["files"]["event_csv"]).stat().st_size > 0
    assert not (folder / seg["files"]["event_xlsx"]).exists()
    assert not list(folder.glob("*.partial"))


def test_finishing_backlog_is_reported(tmp_path, config, monkeypatch):
    import threading as th

    release = th.Event()
    monkeypatch.setattr(SessionLog, "_finish_segment", lambda self, seg, entry: release.wait(30))
    warnings = []
    clock = Clock()
    s = SessionLog(tmp_path / "sess", config, SessionSettings(segment_seconds=10), raw_capture=RawHooks(),
                   on_warning=warnings.append, clock=lambda: clock.wall, perf=lambda: clock.perf,
                   disk_free=lambda p: 1 << 40)
    s.start()
    try:
        for _ in range(4):
            clock.advance(11)
            s.tick()
        assert any("segments are still being finished" in w for w in warnings), warnings
    finally:
        release.set()
        s.stop()
        assert s.await_finished(30)


def test_segment_that_cannot_open_is_recorded_and_raised(tmp_path, config, monkeypatch):
    """task.md Phase 4: rotation into a segment that cannot be created
    (permission, disk, lock). Everything already written stays intact, the
    failure is raised (the GUI stops logging and says so) and session.json
    records why the recording ended."""
    import app.serial_logging.session as sm

    clock = Clock()
    s, _ = _session(tmp_path, config, clock, segment_seconds=10, fixed_step_s=1.0)
    s.start()
    for f in _frames(config):
        s.log_frame(f, clock.perf)

    def denied(self, *a, **k):
        raise PermissionError(13, "Access is denied")

    monkeypatch.setattr(sm.CsvSegmentWriter, "__init__", denied)
    clock.advance(11)
    with pytest.raises(PermissionError):
        s.tick()
    monkeypatch.undo()
    assert s.segment_index is None
    s.stop()
    assert s.await_finished(60)
    m = read_manifest(tmp_path / "sess")
    assert [x["status"] for x in m["segments"]] == ["closed"]
    assert (tmp_path / "sess" / m["segments"][0]["files"]["event_csv"]).stat().st_size > 0
    assert any("segment 2 could not be opened" in e["text"] for e in m["events"]), m["events"]


def test_gui_stops_logging_with_a_message_when_housekeeping_fails():
    from unittest.mock import MagicMock

    from app.ui.logging_session import LoggingSessionMixin

    class Host(LoggingSessionMixin):
        pass

    host = Host()
    host._session_log = MagicMock()
    host._session_log.tick.side_effect = PermissionError("Access is denied")
    host._on_logger_error = MagicMock()
    host._on_session_tick()
    host._on_logger_error.assert_called_once()
    assert "Access is denied" in host._on_logger_error.call_args.args[0]


def test_recovery_leaves_a_closed_session_untouched(tmp_path, config):
    import hashlib

    from app.serial_logging.session import recover_session

    clock = Clock()
    s, _ = _session(tmp_path, config, clock, segment_seconds=10)
    s.start()
    for f in _frames(config):
        s.log_frame(f, clock.perf)
    s.stop()
    assert s.await_finished(60)
    folder = tmp_path / "sess"
    before = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in folder.iterdir() if p.is_file()}
    assert recover_session(folder) is False
    after = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in folder.iterdir() if p.is_file()}
    assert after == before and read_manifest(folder)["status"] == "closed"


def test_crash_during_rotation_is_recovered(tmp_path, config, monkeypatch):
    """The process dies between closing a segment and opening the next one
    (manifest: last segment 'closing', session 'recording')."""
    import app.serial_logging.session as sm
    from app.serial_logging.session import recover_session

    clock = Clock()
    s, _ = _session(tmp_path, config, clock, segment_seconds=10, fixed_step_s=1.0)
    s.start()
    for f in _frames(config) * 5:
        s.log_frame(f, clock.perf)
        clock.advance(0.5)
    closed = []
    # Dies after the writers closed but before finishing / opening the next
    # segment. (A real crash cannot race recovery against live writer
    # threads, so let them complete first.)
    monkeypatch.setattr(SessionLog, "_finish_segment", lambda self, seg, entry: closed.append(seg))
    monkeypatch.setattr(sm.SessionLog, "_open_segment", lambda self, *a, **k: None)
    clock.advance(11)
    s.tick()
    for seg in closed:
        seg.event.await_drain(None)
        seg.fixed.join()
    folder = tmp_path / "sess"
    m = read_manifest(folder)
    assert m["status"] == "recording" and m["segments"][-1]["status"] == "closing"
    monkeypatch.undo()
    assert recover_session(folder)
    m = read_manifest(folder)
    assert m["status"] == "recovered"
    assert all(x["status"] in ("closed", "recovered") for x in m["segments"])
    assert (folder / m["segments"][0]["files"]["event_csv"]).stat().st_size > 0


def test_disk_full_while_writing_is_reported_and_recorded(tmp_path, config, monkeypatch):
    """A write fails mid-segment (ENOSPC): the session's error callback
    fires (the GUI stops logging and says so) and the segment is
    closed_with_errors with the reason in session.json."""
    import csv as _csv
    import time as _t

    import app.serial_logging.decoded_logger as dl

    real_writer = _csv.writer

    class FullDisk:
        def __init__(self, fp):
            self._w = real_writer(fp)
            self.n = 0

        def writerow(self, row):
            self.n += 1
            if self.n > 3:                      # header + 2 rows fit, then the disk is full
                raise OSError(28, "No space left on device")
            return self._w.writerow(row)

    monkeypatch.setattr(dl.csv, "writer", FullDisk)
    errors = []
    clock = Clock()
    s = SessionLog(tmp_path / "sess", config, SessionSettings(segment_seconds=3600), raw_capture=RawHooks(),
                   on_error=errors.append, clock=lambda: clock.wall, perf=lambda: clock.perf,
                   disk_free=lambda p: 1 << 40)
    s.start()
    deadline = _t.monotonic() + 10
    while not errors and _t.monotonic() < deadline:     # keep logging until the writer error surfaces
        for f in _frames(config):
            s.log_frame(f, clock.perf)
        clock.advance(0.1)
        _t.sleep(0.02)
    assert errors and "No space left" in errors[0]
    s.stop()
    assert s.await_finished(60)
    seg = read_manifest(tmp_path / "sess")["segments"][0]
    assert seg["status"] == "closed_with_errors"
    assert any("No space left" in e for e in seg["errors"]), seg["errors"]


def test_fixed_step_write_failure_is_reported_at_the_next_tick(tmp_path, config):
    """The fixed-step writer thread dies on a write error (disk full): the
    error is reported at the next tick, not an hour later at rotation, and
    rows are no longer queued into a writer that cannot write them."""
    import time as _t

    errors = []
    clock = Clock()
    s = SessionLog(tmp_path / "sess", config, SessionSettings(segment_seconds=3600, fixed_step_s=0.5),
                   raw_capture=RawHooks(), on_error=errors.append, clock=lambda: clock.wall,
                   perf=lambda: clock.perf, disk_free=lambda p: 1 << 40)
    s.start()
    fixed = s._segment.fixed
    real = fixed._writer

    class FullDisk:
        def writerow(self, row):
            if fixed.rows_written >= 2:
                raise OSError(28, "No space left on device")
            return real.writerow(row)

    fixed._writer = FullDisk()
    deadline = _t.monotonic() + 10
    while fixed.error is None and _t.monotonic() < deadline:
        for f in _frames(config):
            s.log_frame(f, clock.perf)
        clock.advance(0.5)
        _t.sleep(0.02)
    assert fixed.error and "No space left" in fixed.error
    s.tick()
    assert errors and "No space left" in errors[0]
    s.tick()
    assert len(errors) == 1                          # reported once
    queued = fixed._queue.qsize()
    for _ in range(50):
        fixed.write(["x"])
    assert fixed._queue.qsize() == queued and fixed.dropped >= 50
    s.stop()
    assert s.await_finished(60)
    seg = read_manifest(tmp_path / "sess")["segments"][0]
    assert seg["status"] == "closed_with_errors"
