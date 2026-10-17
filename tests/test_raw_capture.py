"""Raw wire capture (.bhraw) and offline decoded-log rebuild.

The raw capture is the lossless record of a logging session, so the key
guarantees are: every read lands in the file exactly once, a crash-truncated
tail is tolerated, and rebuilding the decoded log from the capture gives the
same rows the live logger wrote.
"""

from __future__ import annotations

import socket
import threading
import time
from pathlib import Path

import pytest
from openpyxl import load_workbook
from PySide6.QtCore import Qt

from app.decoder.config_loader import load_config
from app.protocol.packet_builder import build_packet
from app.serial_io.serial_worker import PollingWorker, SerialSettings
from app.serial_logging.decoded_logger import DecodedLogger
from app.serial_logging.raw_capture import (
    RX, TX, RawCapture, iter_raw_capture, read_metadata, rebuild_decoded_log,
)

TEMPLATE = Path(__file__).resolve().parents[1] / "app" / "resources" / "config_template"


def test_round_trip_preserves_every_record(tmp_path):
    path = tmp_path / "s.bhraw"
    cap = RawCapture(path, {"log_started_perf": 1.5, "note": "x"})
    records = [(RX if i % 3 else TX, 1.0 + i / 1000, bytes([i % 256]) * (1 + i % 40)) for i in range(5000)]
    for r in records:
        cap.write(*r)
    cap.close()

    assert read_metadata(path) == {"log_started_perf": 1.5, "note": "x"}
    assert list(iter_raw_capture(path)) == records
    assert cap.error is None


def test_truncated_tail_is_ignored(tmp_path):
    path = tmp_path / "t.bhraw"
    cap = RawCapture(path)
    cap.write(RX, 1.0, b"\x01\x02\x03")
    cap.write(RX, 2.0, b"\x04\x05\x06")
    cap.close()
    data = path.read_bytes()
    path.write_bytes(data[:-2])  # simulate a crash mid-record

    assert list(iter_raw_capture(path)) == [(RX, 1.0, b"\x01\x02\x03")]


def test_writes_after_close_are_ignored(tmp_path):
    path = tmp_path / "c.bhraw"
    cap = RawCapture(path)
    cap.write(RX, 1.0, b"a")
    cap.close()
    cap.write(RX, 2.0, b"b")
    cap.close()
    assert list(iter_raw_capture(path)) == [(RX, 1.0, b"a")]


def test_rejects_non_capture_file(tmp_path):
    bogus = tmp_path / "x.bhraw"
    bogus.write_bytes(b"not a capture")
    with pytest.raises(ValueError):
        read_metadata(bogus)


def _data_rows(xlsx: Path) -> list:
    wb = load_workbook(xlsx, read_only=True)
    try:
        return [list(r) for r in wb["Data"].iter_rows(values_only=True)]
    finally:
        wb.close()


def test_rebuild_matches_live_decoded_log(tmp_path):
    """Stream frames over TCP into a real PollingWorker while logging live,
    then rebuild from the raw capture: the Data sheets must be identical."""
    config = load_config(TEMPLATE)
    frames = [
        build_packet(config.protocol, fid, bytes((i * 7 + fid) & 0xFF for i in range(f.payload_length)))
        for fid, f in config.frames.items()
    ]

    srv = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    srv.bind(("127.0.0.1", 0))
    srv.listen(1)
    port = srv.getsockname()[1]

    def device() -> None:
        conn, _ = srv.accept()
        # Split frames across sends so the parser sees partial frames too.
        blob = b"".join(frames[i % len(frames)] for i in range(600))
        for i in range(0, len(blob), 37):
            conn.sendall(blob[i:i + 37])
            time.sleep(0.0005)
        time.sleep(0.3)
        conn.close()

    t = threading.Thread(target=device, daemon=True)
    t.start()

    live_path = tmp_path / "live.xlsx"
    raw_path = tmp_path / "live.bhraw"
    log_start = time.perf_counter()
    live = DecodedLogger(live_path, config)
    live.open()

    worker = PollingWorker(
        SerialSettings(port="127.0.0.1", connection_type="tcp", host="127.0.0.1", port_num=port),
        config.protocol, [], decode_config=config,
    )
    count = {"n": 0}

    def on_batch(batch) -> None:
        for packet, decoded, rx_perf in batch:
            if packet.ok and decoded is not None:
                live.log_frame(decoded, int((rx_perf - log_start) * 1000))
                count["n"] += 1

    worker.packets_received.connect(on_batch, Qt.ConnectionType.DirectConnection)
    worker.start_raw_capture(raw_path, {"log_started_perf": log_start})
    worker.open()
    deadline = time.monotonic() + 10
    while count["n"] < 600 and time.monotonic() < deadline:
        time.sleep(0.05)
    worker.close()
    live.close()
    t.join(timeout=2)
    srv.close()

    assert count["n"] == 600
    rebuilt_path = tmp_path / "rebuilt.xlsx"
    assert rebuild_decoded_log(raw_path, config, rebuilt_path) == 600
    live_rows = _data_rows(live_path)
    assert len(live_rows) > 100
    assert _data_rows(rebuilt_path) == live_rows


def test_worker_rotation_keeps_every_byte_in_exactly_one_file(tmp_path):
    """Rotating while the worker thread records must not drop or duplicate
    a chunk (the swap and each record are serialised)."""
    import threading

    from app.serial_io.serial_worker import PollingWorker, SerialSettings
    from app.serial_logging.raw_capture import iter_raw_capture
    from tests.conftest import dummy_protocol_config

    class Port:
        is_open = True

        def read(self, n):
            return b"\x5A" * n

    worker = PollingWorker(SerialSettings(port="x"), dummy_protocol_config(), [])
    worker._serial = Port()
    worker.start_raw_capture(tmp_path / "r0.bhraw")
    stop = threading.Event()
    reads = [0]

    def reader():
        while not stop.is_set():
            worker._read_rx(7)
            reads[0] += 1

    t = threading.Thread(target=reader)
    t.start()
    for i in range(1, 30):
        assert worker.rotate_raw_capture(tmp_path / f"r{i}.bhraw") is None
    stop.set()
    t.join()
    worker.stop_raw_capture()
    worker._serial = None
    total = sum(len(d) for i in range(30) for _dir, _t, d in iter_raw_capture(tmp_path / f"r{i}.bhraw"))
    assert total == 7 * reads[0]
